"""Shared python-docx helpers for WP6 (and session-1's editorial scripts).

Document conventions observed during inspection of Project_Report_Draft.docx:
  - table cells: 10 pt (127000 EMU) body text, bold header row
  - tables: single borders, sz=4, color auto, on all six edges
  - code blocks: Consolas 9 pt (114300 EMU), one paragraph per line,
    wrapped in ``` fence paragraphs
  - figures: centered paragraph, picture scaled to a fixed width
  - the document uses python-docx's default Normal-based body text
"""

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _el(p_or_tbl):
    """Return the underlying lxml element of a paragraph or table."""
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    if isinstance(p_or_tbl, Paragraph):
        return p_or_tbl._p
    if isinstance(p_or_tbl, Table):
        return p_or_tbl._tbl
    raise TypeError(f"expected Paragraph or Table, got {type(p_or_tbl)}")


def move_after(anchor, element):
    """Place `element` (a Paragraph/Table or their lxml element) immediately after `anchor`."""
    anchor_el = _el(anchor) if not hasattr(anchor, 'tag') else anchor
    anchor_el.addnext(_el(element) if not hasattr(element, 'tag') else element)


def _make_paragraph(doc, text='', style=None, align=None, bold=False,
                    size=None, mono=False):
    p = doc.add_paragraph()
    if style:
        try:
            p.style = doc.styles[style]
        except KeyError:
            # styles part in this document resolves by id but not by name;
            # set w:pStyle directly, translating to the built-in id
            style_id = style.replace(' ', '')
            from docx.oxml import OxmlElement as _OX
            pPr = p._p.get_or_add_pPr()
            pStyle = _OX('w:pStyle')
            pStyle.set(qn('w:val'), style_id)
            pPr.insert(0, pStyle)
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        if bold:
            run.font.bold = True
        if size is not None:
            run.font.size = Pt(size)
        if mono:
            run.font.name = 'Consolas'
    return p


def new_para_after(doc, anchor, text, **kw):
    p = _make_paragraph(doc, text, **kw)
    move_after(anchor, p)
    return p


def new_heading_after(doc, anchor, text, level=2):
    p = _make_paragraph(doc, text, style=f'Heading {level}')
    move_after(anchor, p)
    return p


def find_para(doc, prefix):
    for p in doc.paragraphs:
        if p.text.startswith(prefix):
            return p
    raise LookupError(f"no paragraph starts with: {prefix[:60]!r}")


def _borders_xml():
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), 'auto')
        b.append(e)
    return b


def fill_cell(cell, text, bold=False, size=10, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.bold = bold
    r.font.size = Pt(size)
    if align is not None:
        p.alignment = align


def new_table_after(doc, anchor, headers, rows, bold_header=True,
                    cell_size=10, align_header=WD_ALIGN_PARAGRAPH.CENTER):
    """Insert a bordered table after `anchor`. Returns the Table object."""
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tblPr = t._tbl.tblPr
    tblStyle = tblPr.find(qn('w:tblStyle'))
    if tblStyle is None:
        tblStyle = OxmlElement('w:tblStyle')
        tblPr.append(tblStyle)
    tblStyle.set(qn('w:val'), 'TableGrid')
    existing = tblPr.find(qn('w:tblBorders'))
    if existing is not None:
        tblPr.remove(existing)
    tblPr.append(_borders_xml())
    for j, h in enumerate(headers):
        fill_cell(t.rows[0].cells[j], h, bold=bold_header, size=cell_size,
                  align=align_header)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            fill_cell(t.rows[i].cells[j], str(val), size=cell_size)
    move_after(anchor, t)
    return t


def delete_table(tbl):
    el = _el(tbl)
    el.getparent().remove(el)


def set_para_text(para, new_text):
    """Replace paragraph body with a single run, reusing the first run's rPr."""
    keep = None
    for r in para.runs:
        keep = r._r.find(qn('w:rPr'))
        break
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    run_el = OxmlElement('w:r')
    if keep is not None:
        run_el.append(keep)
    t = OxmlElement('w:t')
    t.text = new_text
    t.set(qn('xml:space'), 'preserve')
    run_el.append(t)
    para._p.append(run_el)


def new_code_block_after(doc, anchor, code):
    lines = [ln.rstrip() for ln in code.strip('\n').split('\n')]
    p0 = _make_paragraph(doc, '```', mono=True, size=9)
    move_after(anchor, p0)
    last = p0
    for ln in lines:
        last = new_para_after(doc, last, ln, mono=True, size=9)
    last = new_para_after(doc, last, '```', mono=True, size=9)
    return last


def insert_image_after(doc, anchor, path, width_in=6.0):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(path, width=Inches(width_in))
    move_after(anchor, p)
    return p


def dump_structure(doc, path):
    lines = []
    for i, p in enumerate(doc.paragraphs):
        style = p.style.name if p.style else 'None'
        lines.append(f"P{str(i).rjust(4)} [{style}] {p.text}")
    for ti, tbl in enumerate(doc.tables):
        lines.append(f"TABLE {ti}: {len(tbl.rows)}x{len(tbl.columns)}")
    with open(path, 'w') as f:
        f.write("\n".join(lines))
    return len(doc.paragraphs), len(doc.tables)